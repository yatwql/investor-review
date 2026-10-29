"""持仓文件在线录入编辑层 —— 「读 → 草稿 → 原子写回」单源（TUI / WEB 共用）。

职责边界：
  - 只做**持仓文件本身的读写语义**（账户 = 工作表名、行 = 名称/代码/份额/每份成本），
    不做交互、不取网络数据（净值由调用方经 :mod:`src.python.fetcher.nav_lookup` 取得后
    作为成交价传入）。
  - 草稿语义：:class:`HoldingsBook` 是内存工作副本，``apply`` 只改内存并记录待保存意图；
    ``save`` 才写文件。取消 / 退出面板不触达文件即天然无副作用。

写回契约：
  - 主表四列就地重写（表头保留），账户顺序沿用工作簿内工作表顺序，新建账户追加在末尾；
  - 非持仓页签（交易流水 / 分红流水 / 任意自定义页签）原样保留；
  - 落盘经 :func:`core.atomic_write.write_bytes_atomic`（同目录临时文件 + ``os.replace``），
    不存在半写状态。

成本口径（规则的事实位在实现计划表「持仓录入」条）：买入按份额加权平均成本；卖出
保留原每份成本、只减份额，份额归零时移除该行。
"""

from __future__ import annotations

import io
import logging
import os
from dataclasses import dataclass, field

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from src.python.core.atomic_write import write_bytes_atomic
from src.python.core.code_symbols import strip_to_digits
from src.python.core.code_utils import (
    is_a_share_stock,
    is_exchange_fund_code,
    is_hk_stock_code,
    is_otc_fund_by_name,
)
from src.python.core.models import Holding
from src.python.core.reader import (
    _DIVIDEND_SHEET_NAME,
    _EXPECTED_HEADER,
    _TRADE_SHEET_NAME,
    _match_header,
    read_holdings,
)

logger = logging.getLogger("invest")

SIDE_BUY = "buy"
SIDE_SELL = "sell"
SIDE_LABELS: dict[str, str] = {SIDE_BUY: "买入", SIDE_SELL: "卖出"}

PRICE_KIND_EXCHANGE = "exchange"
PRICE_KIND_OTC = "otc"

# 份额/成本落盘精度：基金份额 2 位、净值 4 位；留 4/6 位吸收加权平均的除不尽余数
_SHARES_DIGITS = 4
_COST_DIGITS = 6
# 份额归零判据（浮点比较容差，远小于 0.01 份的最小交易单位）
_ZERO_EPS = 1e-9
# 回读校验容差（仅吸收浮点表示误差，不掩盖任何实质性差异）
_VERIFY_EPS = 1e-6


@dataclass(frozen=True)
class TradeIntent:
    """一笔待保存的买卖意图（P/B 动作产物，S 动作的输入）。

    Attributes:
        account:    账户名（= 工作表名）
        name:       证券/基金名称
        code:       证券代码
        side:       ``SIDE_BUY`` 或 ``SIDE_SELL``
        shares:     成交份额
        price:      成交价 / 单位净值（场外为交易日官方净值）
        trade_date: 成交日期（YYYY-MM-DD；场外必填，场内可空）
    """

    account: str
    name: str
    code: str
    side: str
    shares: float
    price: float
    trade_date: str = ""

    @property
    def action_label(self) -> str:
        """中文操作方向（买入 / 卖出）。"""
        return SIDE_LABELS.get(self.side, self.side)


@dataclass
class HoldingsBook:
    """持仓文件的内存工作副本 + 草稿清单。

    Attributes:
        filepath:    源文件路径（``save`` 的写入目标）
        accounts:    账户 → 行列表（dict 保插入序 = 工作簿内工作表顺序）
        drafts:      已应用但尚未保存的买卖意图（按录入顺序）
        dirty:       被本次录入改动过的 (账户, 归一代码) 集合，供交互层标注「未保存」
    """

    filepath: str
    accounts: dict[str, list[Holding]] = field(default_factory=dict)
    drafts: list[TradeIntent] = field(default_factory=list)
    dirty: set[tuple[str, str]] = field(default_factory=set)

    @property
    def account_names(self) -> list[str]:
        """账户名列表（工作簿顺序）。"""
        return list(self.accounts)

    def rows(self, account: str) -> list[Holding]:
        """账户内持仓行（未知账户返回空列表）。"""
        return self.accounts.get(account, [])

    def is_dirty(self, account: str, code: str) -> bool:
        """该账户下的某代码是否被本次未保存的录入改动过（代码归一后比较）。"""
        return ((account or "").strip(), strip_to_digits(code)) in self.dirty

    def dirty_accounts(self) -> set[str]:
        """被本次未保存录入改动过的账户集合。"""
        return {account for account, _code in self.dirty}

    def total_rows(self) -> int:
        """全部账户的持仓行数。"""
        return sum(len(rows) for rows in self.accounts.values())

    def apply(self, intent: TradeIntent) -> None:
        """把一笔买卖意图应用到内存副本（不落盘）。

        Args:
            intent: 买卖意图（已在交互层完成字段校验）。

        Raises:
            ValueError: 方向非法 / 账户为空 / 卖出超份额 / 卖出无此持仓。
        """
        if intent.side not in SIDE_LABELS:
            raise ValueError(f"未知识别操作方向: {intent.side}")
        account = (intent.account or "").strip()
        if not account:
            raise ValueError("账户名不能为空")
        if intent.shares <= 0:
            raise ValueError("份额必须大于 0")
        if intent.price <= 0:
            raise ValueError("价格必须大于 0")

        rows = self.accounts.setdefault(account, [])
        existing = _find_row(rows, intent.code)

        if intent.side == SIDE_BUY:
            if existing is None:
                rows.append(
                    Holding(
                        account=account,
                        name=intent.name,
                        code=intent.code,
                        shares=round(intent.shares, _SHARES_DIGITS),
                        cost_price=round(intent.price, _COST_DIGITS),
                    )
                )
            else:
                new_shares = existing.shares + intent.shares
                weighted = (existing.shares * existing.cost_price + intent.shares * intent.price) / new_shares
                existing.shares = round(new_shares, _SHARES_DIGITS)
                existing.cost_price = round(weighted, _COST_DIGITS)
        else:
            if existing is None:
                raise ValueError(f"账户「{account}」中没有 {intent.name}({intent.code}) 的持仓，无法卖出")
            if intent.shares > existing.shares + _ZERO_EPS:
                raise ValueError(
                    f"卖出份额 {intent.shares:g} 超过当前持有 {existing.shares:g}（{intent.name}({intent.code})）"
                )
            remaining = existing.shares - intent.shares
            if remaining <= _ZERO_EPS:
                rows.remove(existing)
            else:
                existing.shares = round(remaining, _SHARES_DIGITS)

        self.drafts.append(intent)
        self.dirty.add((account, strip_to_digits(intent.code)))
        logger.info(
            "[持仓录入] 草稿 %s %s(%s) 份额=%s 价格=%s 账户=%s",
            intent.action_label,
            intent.name,
            intent.code,
            intent.shares,
            intent.price,
            account,
        )

    def revert_last(self) -> TradeIntent | None:
        """撤销最后一笔录入（从文件重读后重放剩余变更，保证状态与文件一致）。

        Returns:
            被撤销的意图；无可撤销项返回 ``None``。
        """
        if not self.drafts:
            return None
        reverted = self.drafts.pop()
        self._replay(list(self.drafts))
        return reverted

    def _replay(self, intents: list[TradeIntent]) -> None:
        """以文件为基线重放给定变更（撤销后的重建路径，不落盘）。"""
        fresh = load_book(self.filepath)
        self.accounts = {account: list(rows) for account, rows in fresh.accounts.items()}
        self.drafts = []
        self.dirty = set()
        for intent in intents:
            self.apply(intent)

    def save(self, *, append_flows: bool = False) -> None:
        """把内存副本原子写回持仓文件（含可选交易流水追加），回读校验后清空变更。

        Args:
            append_flows: 是否把本次草稿追加进「交易流水」页签。

        Raises:
            ValueError: 页签结构冲突（新账户名被非持仓页签占用 / 交易流水表头不匹配），
                或写盘后回读校验不一致。
            OSError: 落盘失败。
        """
        data = build_workbook_bytes(self, append_flows=append_flows)
        if not write_bytes_atomic(self.filepath, data, log_tag="holdings", noun="持仓文件"):
            raise OSError(f"持仓文件写入失败: {self.filepath}")
        self._verify_written()
        logger.info("[持仓录入] 已保存并回读校验 %d 笔变更 → %s", len(self.drafts), self.filepath)
        self.drafts.clear()
        self.dirty.clear()

    def _verify_written(self) -> None:
        """写盘后回读校验：内存结果与文件解析结果必须逐行一致（份额/成本）。

        Raises:
            ValueError: 回读结果与内存不一致（草稿保留，用户可重试或核对文件）。
        """
        expected = {
            (account, strip_to_digits(row.code)): row for account, rows in self.accounts.items() for row in rows
        }
        actual = {(row.account.strip(), strip_to_digits(row.code)): row for row in read_holdings(self.filepath)}
        problems: list[str] = []
        for key in sorted(set(expected) - set(actual)):
            problems.append(f"{key[0]}/{key[1]} 丢失")
        for key in sorted(set(actual) - set(expected)):
            problems.append(f"{key[0]}/{key[1]} 多出")
        for key in sorted(set(expected) & set(actual)):
            want, got = expected[key], actual[key]
            if abs(want.shares - got.shares) > _VERIFY_EPS or abs(want.cost_price - got.cost_price) > _VERIFY_EPS:
                problems.append(
                    f"{key[0]}/{key[1]} 份额 {got.shares:g}≠{want.shares:g} 或成本 {got.cost_price:g}≠{want.cost_price:g}"
                )
        if problems:
            raise ValueError("写入后回读校验未通过：" + "；".join(problems[:5]))


def price_kind(code: str, name: str = "") -> str:
    """品种价格形态：场内（须录成交价）或场外（按交易日取官方净值）。

    Args:
        code: 6 位证券代码。
        name: 名称（用于 00 前缀重叠区区分 A 股个股与场外基金）。

    Returns:
        ``PRICE_KIND_EXCHANGE`` 或 ``PRICE_KIND_OTC``。
    """
    raw = strip_to_digits(code)
    if not raw:
        return PRICE_KIND_OTC
    if is_a_share_stock(name, raw) or is_exchange_fund_code(raw) or is_hk_stock_code(raw):
        return PRICE_KIND_EXCHANGE
    return PRICE_KIND_OTC


def ambiguous_price_kind(code: str, name: str = "") -> bool:
    """00 代码重叠区且名称无基金特征词时，场内/场外无法仅凭代码判定。

    00 前缀同时覆盖深市主板/中小板与场外基金（005827、007301 等），
    :func:`price_kind` 会按「A 股」处理；本函数把这类**歧义**单独暴露出来，
    由交互层用行情源反推或向用户确认，避免把场外基金当股票要求录成交价。

    Args:
        code: 证券代码。
        name: 名称（用于识别基金特征词）。

    Returns:
        True 表示需一次显式确认。
    """
    raw = strip_to_digits(code)
    if len(raw) != 6 or not raw.startswith("00"):
        return False
    if is_exchange_fund_code(raw) or is_hk_stock_code(raw):
        return False
    return not is_otc_fund_by_name(name, raw)


def price_kind_from_quote(data: dict | None) -> str | None:
    """按行情取数结果反推品种价格形态（歧义区消歧用）。

    行情链路已按代码类型路由：命中场外净值源（``eastmoney`` / ``sina_fund``）即场外；
    命中行情源（``tencent`` / ``sina`` 等）即场内。无数据返回 ``None``（交用户确认）。

    Args:
        data: ``fetch_market_data`` 的返回字典（可为 None）。

    Returns:
        ``PRICE_KIND_OTC`` / ``PRICE_KIND_EXCHANGE`` / ``None``。
    """
    if not data:
        return None
    source_api = str(data.get("source_api") or "").strip()
    if not source_api:
        return None
    if source_api in ("eastmoney", "sina_fund"):
        return PRICE_KIND_OTC
    return PRICE_KIND_EXCHANGE


def load_book(filepath: str) -> HoldingsBook:
    """读取持仓文件构造内存副本。

    Args:
        filepath: 持仓 .xlsx 路径。

    Returns:
        HoldingsBook（账户顺序 = 工作簿内持仓页签顺序）。

    Raises:
        FileNotFoundError: 文件不存在。
        ValueError: 文件格式非法。
    """
    holdings = read_holdings(filepath)
    accounts: dict[str, list[Holding]] = {}
    wb = openpyxl.load_workbook(filepath, read_only=True)
    try:
        for sheet_name in wb.sheetnames:
            if _is_flow_sheet_name(sheet_name):
                continue
            header = [
                str(cell.value or "").strip() for cell in next(wb[sheet_name].iter_rows(min_row=1, max_row=1), ())
            ]
            if _match_header(header, _EXPECTED_HEADER):
                accounts.setdefault(sheet_name.strip(), [])
    finally:
        wb.close()
    for holding in holdings:
        accounts.setdefault(holding.account, []).append(holding)
    return HoldingsBook(filepath=filepath, accounts=accounts)


def build_workbook_bytes(book: HoldingsBook, *, append_flows: bool = False) -> bytes:
    """按内存副本生成完整工作簿字节（不落盘，供写回与测试复用）。

    Args:
        book: 内存副本。
        append_flows: 是否把草稿追加进「交易流水」页签。

    Returns:
        完整 .xlsx 内容。

    Raises:
        ValueError: 新账户名与既有非持仓页签同名 / 交易流水表头不匹配。
    """
    if not os.path.exists(book.filepath):
        raise FileNotFoundError(f"持仓文件不存在: {book.filepath}")
    wb = openpyxl.load_workbook(book.filepath)
    try:
        holding_sheets: dict[str, Worksheet] = {}
        other_titles: set[str] = set()
        for ws in wb.worksheets:
            title = ws.title.strip()
            if _is_holding_sheet(ws):
                holding_sheets[title] = ws
            else:
                other_titles.add(title)

        for account, rows in book.accounts.items():
            ws = holding_sheets.get(account)
            if ws is None:
                if account in other_titles:
                    raise ValueError(f"账户「{account}」与持仓文件中已有的非持仓页签同名，请换一个账户名")
                ws = wb.create_sheet(title=account)
                ws.append(list(_EXPECTED_HEADER))
            elif ws.max_row is not None and ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            for row in rows:
                ws.append(
                    [
                        row.name,
                        row.code,
                        float(round(row.shares, _SHARES_DIGITS)),
                        float(round(row.cost_price, _COST_DIGITS)),
                    ]
                )

        if append_flows and book.drafts:
            _append_flow_rows(wb, book.drafts)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    finally:
        wb.close()


def _append_flow_rows(wb, drafts: list[TradeIntent]) -> None:
    """向「交易流水」页签追加草稿行（缺页签时新建）。

    Raises:
        ValueError: 既有页签表头与持仓文件读取契约（日期/代码/操作/份额/价格）不匹配。
    """
    ws = wb[_TRADE_SHEET_NAME] if _TRADE_SHEET_NAME in wb.sheetnames else None
    has_fee = False
    if ws is None:
        ws = wb.create_sheet(title=_TRADE_SHEET_NAME)
        ws.append(["日期", "代码", "操作", "份额", "价格"])
    else:
        header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if header[:5] != ["日期", "代码", "操作", "份额", "价格"]:
            raise ValueError(f"「{_TRADE_SHEET_NAME}」页签表头不匹配，无法追加记录（应为 日期/代码/操作/份额/价格）")
        has_fee = len(header) >= 6 and header[5] == "费用"
    for intent in drafts:
        row = [intent.trade_date, intent.code, intent.action_label, float(intent.shares), float(intent.price)]
        if has_fee:
            row.append(0.0)
        ws.append(row)


def _find_row(rows: list[Holding], code: str) -> Holding | None:
    """账户内按代码查行（代码归一后比较，容忍前导零/交易所前缀差异）。"""
    target = strip_to_digits(code)
    for row in rows:
        if strip_to_digits(row.code) == target:
            return row
    return None


def _is_flow_sheet_name(sheet_name: str) -> bool:
    """是否为流水类页签（非账户主表）。"""
    return sheet_name.strip() in (_TRADE_SHEET_NAME, _DIVIDEND_SHEET_NAME)


def _is_holding_sheet(ws) -> bool:
    """工作表是否为持仓主表（表头须为 名称/代码/持仓份额/每份成本）。"""
    if _is_flow_sheet_name(ws.title):
        return False
    if ws.max_row is None:
        return False
    header = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    return _match_header(header, _EXPECTED_HEADER)
