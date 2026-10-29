"""持仓录入相关测试共用的持仓 xlsx 构造器（TUI 面板 / 编辑层 / 边缘三处共用）。

原先三个测试文件各自维护一份 `_make_holdings_file`（账户名、表头、流水页签逐字复制），
改一个表头或加一列就得改三处，容易悄悄漂移。收敛到本模块后，默认值只在这里定义，
各用例按需覆写行列或关掉某个页签。
"""

from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

import openpyxl

HOLDINGS_HEADER: list[str] = ["名称", "代码", "持仓份额", "每份成本"]
TRADE_HEADER: list[str] = ["日期", "代码", "操作", "份额", "价格"]

#: 默认股票账户行（两行：一只 A 股 + 一只场外基金，覆盖两类价格形态）
DEFAULT_STOCK_ROWS: list[list[object]] = [
    ["长江电力", "600900", 1000, 20.0],
    ["华夏成长", "007301", 500, 1.5],
]
#: 默认基金账户行（名称含基金特征词，避免落到 00 代码歧义区）
DEFAULT_FUND_ROWS: list[list[object]] = [["易方达蓝筹精选混合", "005827", 200, 2.0]]
#: 默认交易流水行（写回时「保留其他页签」的观察对象）
DEFAULT_FLOW_ROWS: list[list[object]] = [["2026-06-01", "600900", "买入", 100, 19.0]]


def make_holdings_file(
    tmp_path: Path,
    *,
    stock_rows: Sequence[Sequence[object]] | None = None,
    fund_rows: Sequence[Sequence[object]] | None = None,
    flow_rows: Sequence[Sequence[object]] | None = None,
    with_fund_sheet: bool = True,
    with_flow_sheet: bool = True,
    extra_sheets: Mapping[str, Sequence[Sequence[object]]] | None = None,
    filename: str = "持仓.xlsx",
) -> Path:
    """生成一份持仓工作簿（账户页签 = 账户名，首行固定四列表头）。

    Args:
        tmp_path: pytest 临时目录。
        stock_rows: 「股票账户」数据行；``None`` 用默认两行。
        fund_rows: 「基金账户」数据行；``None`` 用默认一行。
        flow_rows: 「交易流水」数据行；``None`` 用默认一行。
        with_fund_sheet: 是否创建「基金账户」页签。
        with_flow_sheet: 是否创建「交易流水」页签。
        extra_sheets: 额外页签 ``{页签名: 数据行}``（首行按给定行原样写入，用作非持仓页签）。
        filename: 输出文件名。

    Returns:
        生成的 .xlsx 路径。
    """
    wb = openpyxl.Workbook()
    stock = wb.worksheets[0]
    stock.title = "股票账户"
    stock.append(list(HOLDINGS_HEADER))
    for row in DEFAULT_STOCK_ROWS if stock_rows is None else stock_rows:
        stock.append(list(row))

    if with_fund_sheet:
        fund = wb.create_sheet("基金账户")
        fund.append(list(HOLDINGS_HEADER))
        for row in DEFAULT_FUND_ROWS if fund_rows is None else fund_rows:
            fund.append(list(row))

    if with_flow_sheet:
        flow = wb.create_sheet("交易流水")
        flow.append(list(TRADE_HEADER))
        for row in DEFAULT_FLOW_ROWS if flow_rows is None else flow_rows:
            flow.append(list(row))

    for title, rows in (extra_sheets or {}).items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(list(row))

    path = tmp_path / filename
    wb.save(path)
    return path
