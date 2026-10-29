# -*- coding: utf-8 -*-
"""持仓录入编辑层（core/holdings_editor.py）单元测试。

覆盖「读文件 → 内存草稿 → 原子写回」的正面路径：账户/行保序、买入加权平均成本、
卖出保留原成本、写回后其他页签不受影响、可选交易流水追加、新账户建页。
边缘与异常场景见 ``test_holdings_editor_edge.py``。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.python.core.holdings_editor import (
    PRICE_KIND_EXCHANGE,
    PRICE_KIND_OTC,
    SIDE_BUY,
    SIDE_SELL,
    TradeIntent,
    ambiguous_price_kind,
    build_workbook_bytes,
    load_book,
    price_kind,
    resolve_entry_kind,
)
from src.python.core.reader import read_holdings
from src.test._holdings_fixture import make_holdings_file

pytestmark = [pytest.mark.unit, pytest.mark.unit_core]


def test_load_book_keeps_account_order_and_rows(tmp_path: Path):
    book = load_book(str(make_holdings_file(tmp_path)))
    assert book.account_names == ["股票账户", "基金账户"]
    assert book.total_rows() == 3
    assert [row.code for row in book.rows("股票账户")] == ["600900", "007301"]
    assert book.drafts == []


def test_buy_existing_asset_uses_weighted_average_cost(tmp_path: Path):
    book = load_book(str(make_holdings_file(tmp_path)))
    book.apply(
        TradeIntent(
            account="股票账户",
            name="长江电力",
            code="600900",
            side=SIDE_BUY,
            shares=1000,
            price=30.0,
            trade_date="2026-07-01",
        )
    )
    row = next(row for row in book.rows("股票账户") if row.code == "600900")
    assert row.shares == pytest.approx(2000)
    assert row.cost_price == pytest.approx(25.0)


def test_buy_new_asset_appends_row(tmp_path: Path):
    book = load_book(str(make_holdings_file(tmp_path)))
    book.apply(
        TradeIntent(
            account="基金账户",
            name="易方达消费",
            code="110022",
            side=SIDE_BUY,
            shares=300,
            price=1.2345,
            trade_date="2026-07-01",
        )
    )
    rows = book.rows("基金账户")
    assert [row.code for row in rows] == ["005827", "110022"]
    assert rows[-1].shares == pytest.approx(300)


def test_sell_partial_keeps_original_cost(tmp_path: Path):
    book = load_book(str(make_holdings_file(tmp_path)))
    book.apply(
        TradeIntent(
            account="股票账户",
            name="长江电力",
            code="600900",
            side=SIDE_SELL,
            shares=400,
            price=25.0,
            trade_date="2026-07-01",
        )
    )
    row = next(row for row in book.rows("股票账户") if row.code == "600900")
    assert row.shares == pytest.approx(600)
    assert row.cost_price == pytest.approx(20.0)


def test_sell_all_removes_row(tmp_path: Path):
    book = load_book(str(make_holdings_file(tmp_path)))
    book.apply(
        TradeIntent(
            account="基金账户",
            name="易方达蓝筹",
            code="005827",
            side=SIDE_SELL,
            shares=200,
            price=2.5,
            trade_date="2026-07-01",
        )
    )
    assert book.rows("基金账户") == []
    assert book.drafts[-1].side == SIDE_SELL


def test_dirty_tracking_marks_only_touched_rows(tmp_path: Path):
    """未保存标注按 (账户, 归一代码) 记录：只标被本次录入改过的行，保存后清空。"""
    path = make_holdings_file(tmp_path)
    book = load_book(str(path))
    book.apply(
        TradeIntent(
            account="股票账户",
            name="长江电力",
            code="sh600900",  # 带前缀也能命中同一行
            side=SIDE_BUY,
            shares=100,
            price=21.0,
            trade_date="2026-07-01",
        )
    )
    assert book.is_dirty("股票账户", "600900") is True
    assert book.is_dirty("股票账户", "007301") is False
    assert book.dirty_accounts() == {"股票账户"}

    book.save()

    assert book.is_dirty("股票账户", "600900") is False
    assert book.dirty_accounts() == set()


def test_save_writes_rows_and_keeps_other_sheets(tmp_path: Path):
    path = make_holdings_file(tmp_path)
    book = load_book(str(path))
    book.apply(
        TradeIntent(
            account="股票账户",
            name="长江电力",
            code="600900",
            side=SIDE_BUY,
            shares=500,
            price=22.0,
            trade_date="2026-07-01",
        )
    )
    book.save()

    holdings = read_holdings(str(path))
    row = next(item for item in holdings if item.code == "600900")
    assert row.shares == pytest.approx(1500)
    assert row.cost_price == pytest.approx((1000 * 20.0 + 500 * 22.0) / 1500)
    assert book.drafts == []

    wb = openpyxl.load_workbook(path)
    try:
        flow = wb["交易流水"]
        assert flow.max_row == 2  # 未开启追加：流水页签原样保留
        assert flow.cell(row=2, column=2).value == "600900"
    finally:
        wb.close()


def test_save_appends_trade_flow_when_enabled(tmp_path: Path):
    path = make_holdings_file(tmp_path)
    book = load_book(str(path))
    book.apply(
        TradeIntent(
            account="股票账户",
            name="华夏成长",
            code="007301",
            side=SIDE_SELL,
            shares=100,
            price=1.66,
            trade_date="2026-07-02",
        )
    )
    book.save(append_flows=True)

    wb = openpyxl.load_workbook(path)
    try:
        flow = wb["交易流水"]
        assert flow.max_row == 3
        added = [flow.cell(row=3, column=col).value for col in range(1, 6)]
        assert added[:3] == ["2026-07-02", "007301", "卖出"]
        assert added[3] == pytest.approx(100)
        assert added[4] == pytest.approx(1.66)
    finally:
        wb.close()


def test_save_creates_new_account_sheet(tmp_path: Path):
    path = make_holdings_file(tmp_path)
    book = load_book(str(path))
    book.apply(
        TradeIntent(
            account="新账户",
            name="贵州茅台",
            code="600519",
            side=SIDE_BUY,
            shares=10,
            price=1700.0,
            trade_date="2026-07-01",
        )
    )
    book.save()

    holdings = read_holdings(str(path))
    assert any(item.account == "新账户" and item.code == "600519" for item in holdings)


def test_build_workbook_bytes_is_pure(tmp_path: Path):
    path = make_holdings_file(tmp_path)
    before = path.read_bytes()
    book = load_book(str(path))
    book.apply(
        TradeIntent(
            account="股票账户",
            name="长江电力",
            code="600900",
            side=SIDE_BUY,
            shares=100,
            price=21.0,
            trade_date="2026-07-01",
        )
    )
    data = build_workbook_bytes(book)
    assert data[:2] == b"PK"  # 有效 xlsx 容器
    assert path.read_bytes() == before  # 只生成字节，不动文件
    assert len(book.drafts) == 1


@pytest.mark.parametrize(
    ("code", "name", "expected"),
    [
        ("600900", "长江电力", PRICE_KIND_EXCHANGE),
        ("510300", "沪深300ETF", PRICE_KIND_EXCHANGE),
        ("00700", "腾讯控股", PRICE_KIND_OTC),  # 5 位港股码不在 6 位判定域内
        ("007301", "华夏成长混合", PRICE_KIND_OTC),
        ("003096", "中欧医疗健康混合A", PRICE_KIND_OTC),
    ],
)
def test_price_kind_classifies_by_asset_type(code: str, name: str, expected: str):
    assert price_kind(code, name) == expected


def test_revert_last_replays_remaining_changes(tmp_path: Path):
    """撤销最后一笔：以文件为基线重放剩余变更，份额回到撤销前一步的状态。"""
    path = make_holdings_file(tmp_path)
    book = load_book(str(path))
    book.apply(
        TradeIntent(
            account="股票账户",
            name="长江电力",
            code="600900",
            side=SIDE_BUY,
            shares=100,
            price=21.0,
            trade_date="2026-07-01",
        )
    )
    book.apply(
        TradeIntent(
            account="股票账户",
            name="长江电力",
            code="600900",
            side=SIDE_SELL,
            shares=50,
            price=23.0,
            trade_date="2026-07-02",
        )
    )

    reverted = book.revert_last()

    assert reverted is not None and reverted.side == SIDE_SELL
    row = next(item for item in book.rows("股票账户") if item.code == "600900")
    assert row.shares == pytest.approx(1100)
    assert len(book.drafts) == 1

    book.revert_last()

    assert book.drafts == []
    assert book.dirty == set()
    assert book.rows("股票账户")[0].shares == pytest.approx(1000)
    assert book.revert_last() is None


def test_save_verifies_written_file_and_keeps_drafts_on_mismatch(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """回读校验不一致 → 抛错且保留变更（不当作保存成功）。"""
    from src.python.core import holdings_editor

    path = make_holdings_file(tmp_path)
    book = load_book(str(path))
    book.apply(
        TradeIntent(
            account="股票账户",
            name="长江电力",
            code="600900",
            side=SIDE_BUY,
            shares=100,
            price=21.0,
            trade_date="2026-07-01",
        )
    )
    monkeypatch.setattr(holdings_editor, "read_holdings", lambda _path: [])

    with pytest.raises(ValueError, match="回读校验未通过"):
        book.save()

    assert len(book.drafts) == 1


@pytest.mark.parametrize(
    ("code", "name", "ambiguous"),
    [
        ("005827", "易方达蓝筹", True),  # 00 区且名称无基金词 → 需确认
        ("000001", "平安银行", True),  # 同上（股票侧也可能歧义）
        ("005827", "易方达蓝筹精选混合", False),  # 名称含基金词 → 直接场外
        ("600900", "长江电力", False),  # 非 00 区
        ("510300", "沪深300ETF", False),  # 场内基金前缀
    ],
)
def test_ambiguous_price_kind_detects_00_overlap(code: str, name: str, ambiguous: bool):
    assert ambiguous_price_kind(code, name) is ambiguous


@pytest.mark.parametrize(
    ("code", "name", "declared", "expected"),
    [
        ("600900", "长江电力", None, PRICE_KIND_EXCHANGE),
        ("007301", "华夏成长混合", None, "otc"),
        ("005827", "易方达蓝筹", PRICE_KIND_OTC, "otc"),  # 歧义区：显式声明优先
        ("000001", "平安银行", PRICE_KIND_EXCHANGE, PRICE_KIND_EXCHANGE),
    ],
)
def test_resolve_entry_kind_shared_rule(code: str, name: str, declared: str | None, expected: str):
    """品种形态判定单源：显式声明优先，非歧义按代码/名称判定。"""
    assert resolve_entry_kind(code, name, declared) == expected


def test_resolve_entry_kind_rejects_ambiguous_without_declaration():
    """00 重叠区未声明形态 → 抛错交调用方确认（TUI 问用户 / Web 报 400）。"""
    with pytest.raises(ValueError, match="重叠区"):
        resolve_entry_kind("005827", "易方达蓝筹")
