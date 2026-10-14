"""持仓成本区间位置 — 数据契约与渲染层单元测试（report/cost_position + 持仓明细列）。

覆盖：开关门控 / 取数装配 / 降级契约 / 单元格文案 / 页签条件列显隐与脚注。
"""

from __future__ import annotations

from unittest.mock import patch

import pytest
from openpyxl import Workbook

from src.python.core.models import Holding
from src.python.config.features import set_feature_enabled
from src.python.report.cost_position import (
    COST_POSITION_FOOTNOTE,
    compute_cost_position_data,
    cost_position_footnote,
    cost_position_visible,
    format_cost_position_cell,
)
from src.python.report.holdings_detail_sheet import write_holdings_detail_sheet
from src.python.report.market_value import DetailRow

pytestmark = [pytest.mark.unit, pytest.mark.unit_report, pytest.mark.usefixtures("offline_external_sources")]

_CALC = "src.python.report.portfolio_history.PortfolioHistoryCalculator.calculate_for_holding"


def _holding(code: str = "600519", cost: float | None = 1700.0) -> Holding:
    return Holding(account="证券", name=f"品种{code[-2:]}", code=code, shares=100.0, cost_price=cost if cost else 0.0)


def _bars(*closes: float) -> list[dict]:
    return [{"date": f"2026-01-{i + 1:02d}", "close": c} for i, c in enumerate(closes)]


def _detail(code: str = "600519") -> DetailRow:
    return DetailRow(
        account="证券",
        name=f"品种{code[-2:]}",
        code=code,
        price=1800.0,
        yesterday_close=1780.0,
        shares=100.0,
        market_value=180000.0,
        cost=170000.0,
        profit=10000.0,
        profit_rate=0.0588,
        today_profit=2000.0,
        source="腾讯",
        source_api="tencent",
    )


class TestComputeCostPositionData:
    """compute_cost_position_data — 开关门控与契约装配。"""

    def test_switch_off_returns_none(self):
        set_feature_enabled("cost_position", False)
        assert compute_cost_position_data([_holding()], {}) is None

    def test_happy_path_contract(self):
        with patch(_CALC, return_value=_bars(1400.0, 1500.0, 1600.0, 1750.0, 1800.0)):
            data = compute_cost_position_data([_holding()], {})
        assert data is not None and data["available"] is True
        pos = data["positions"]["600519"]
        # 1700 在 [1400,1800] 内 → (1700-1400)/400 = 0.75；两窗口同源同值
        assert pos["short"] == pytest.approx(0.75)
        assert pos["long"] == pytest.approx(0.75)

    def test_fetch_failure_degrades_to_unavailable(self):
        """行情全链失败 → available=False（渲染层静默隐列），不抛错。"""
        with patch(_CALC, return_value=None):
            data = compute_cost_position_data([_holding()], {})
        assert data is not None and data["available"] is False

    def test_empty_holdings_unavailable(self):
        data = compute_cost_position_data([], {})
        assert data is not None and data["available"] is False

    def test_duplicate_codes_fetch_once(self):
        """同代码多账户持仓只取数一次（契约按 code 键控）。"""
        with patch(_CALC, return_value=_bars(100.0, 200.0)) as m:
            compute_cost_position_data([_holding("600519"), _holding("600519")], {})
        assert m.call_count == 1

    def test_missing_cost_allows_other_holdings(self):
        """单品种成本无效不拖垮整体：其余品种有效则 available=True。"""
        with patch(_CALC, return_value=_bars(100.0, 200.0)):
            data = compute_cost_position_data([_holding("600519", cost=None), _holding("000001", cost=150.0)], {})
        assert data is not None and data["available"] is True
        assert data["positions"]["600519"]["short"] is None
        assert data["positions"]["000001"]["short"] is not None


class TestContractHelpers:
    """可见性判据 / 单元格文案 / 脚注单源。"""

    def test_visible_none_or_unavailable(self):
        assert cost_position_visible(None) is False
        assert cost_position_visible({"available": False}) is False
        assert cost_position_visible({"available": True}) is True

    def test_cell_passthrough_and_placeholder(self):
        data = {"available": True, "positions": {"600519": {"short": 0.75, "long": None}}, "footnote": ""}
        assert format_cost_position_cell(data, "600519", "short") == 0.75
        assert format_cost_position_cell(data, "600519", "long") == "--"
        assert format_cost_position_cell(data, "000001", "short") == "--"
        assert format_cost_position_cell(None, "600519", "short") == "--"

    def test_footnote_single_source(self):
        assert cost_position_footnote(None) == COST_POSITION_FOOTNOTE
        data = {"available": True, "positions": {}, "footnote": "X"}
        assert cost_position_footnote(data) == "X"


def _header_row(ws) -> list[str]:
    for row in ws.iter_rows(min_row=1, max_row=5):
        vals = [c.value for c in row if c.value]
        if "名称" in vals:
            return [str(v) for v in vals]
    return []


class TestHoldingsDetailColumns:
    """持仓明细页签「成本区间位置」条件列显隐。"""

    _CONTRACT = {
        "available": True,
        "positions": {"600519": {"short": 0.75, "long": 0.4}},
        "footnote": COST_POSITION_FOOTNOTE,
    }

    def test_contract_absent_keeps_existing_columns(self):
        wb = Workbook()
        write_holdings_detail_sheet(wb.active, [_holding()], [_detail()])
        headers = _header_row(wb.active)
        assert not any("成本区间位置" in h for h in headers)

    def test_unavailable_contract_hides_columns(self):
        wb = Workbook()
        write_holdings_detail_sheet(
            wb.active,
            [_holding()],
            [_detail()],
            cost_position_data={"available": False, "positions": {}, "footnote": ""},
        )
        assert not any("成本区间位置" in h for h in _header_row(wb.active))

    def test_contract_appends_two_columns(self):
        wb = Workbook()
        write_holdings_detail_sheet(wb.active, [_holding()], [_detail()], cost_position_data=self._CONTRACT)
        headers = _header_row(wb.active)
        cost_headers = [h for h in headers if "成本区间位置" in h]
        assert len(cost_headers) == 2

    def test_contract_cells_written(self):
        wb = Workbook()
        ws = wb.active
        write_holdings_detail_sheet(ws, [_holding()], [_detail()], cost_position_data=self._CONTRACT)
        flat = [c.value for row in ws.iter_rows() for c in row]
        assert 0.75 in flat and 0.4 in flat

    def test_missing_code_cell_placeholder(self):
        """契约未覆盖的代码写占位符而非空洞。"""
        contract = {"available": True, "positions": {}, "footnote": COST_POSITION_FOOTNOTE}
        wb = Workbook()
        ws = wb.active
        write_holdings_detail_sheet(ws, [_holding()], [_detail()], cost_position_data=contract)
        flat = [c.value for row in ws.iter_rows() for c in row]
        assert "--" in flat

    def test_footnote_written_when_visible(self):
        wb = Workbook()
        ws = wb.active
        write_holdings_detail_sheet(ws, [_holding()], [_detail()], cost_position_data=self._CONTRACT)
        flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        assert any("成本区间位置" in v and "clamp" in v for v in flat)

    def test_footnote_absent_when_hidden(self):
        wb = Workbook()
        ws = wb.active
        write_holdings_detail_sheet(ws, [_holding()], [_detail()])
        flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        assert not any("clamp" in v for v in flat)
