# -*- coding: utf-8 -*-
"""设计语言契约（docs/managements/DESIGN.md）结构守卫。

契约是 UI 改动的首个读取入口，断言其骨架不漂移：
必需节子集存在、护栏条目连续编号且 DO/DON'T 对偶、角色族覆盖、
文内引用的实现文件真实存在。新增节/新增护栏不破坏断言（子集关系）。
"""

from pathlib import Path

import pytest

pytestmark = [pytest.mark.unit, pytest.mark.unit_report]

REPO_ROOT = Path(__file__).resolve().parents[4]
DESIGN_DOC = REPO_ROOT / "docs" / "managements" / "DESIGN.md"

# 契约必需节（子集断言：只防删除/改名，不锁总数）
REQUIRED_SECTIONS = [
    "Overview",
    "Colors",
    "Typography",
    "Components",
    "Layout",
    "Elevation",
    "Responsive",
    "Data States",
    "Do's and Don'ts",
    "Iteration Guide",
    "Known Gaps",
]

# 颜色角色族必须覆盖的语义域（子集断言）
REQUIRED_COLOR_FAMILIES = [
    "Surface 阶梯",
    "Ink 层级",
    "Border 阶",
    "状态语义",
    "涨跌语义",
    "Chart 专用",
]

# 字阶类别（子集断言）
REQUIRED_TYPE_SCALE = ["章节标题", "正文", "表格", "脚注"]


def _read() -> str:
    assert DESIGN_DOC.exists(), f"设计契约缺失: {DESIGN_DOC}"
    return DESIGN_DOC.read_text(encoding="utf-8")


class TestDesignDocStructure:
    """骨架与内容契约。"""

    def test_required_sections_present(self):
        text = _read()
        headings = [line for line in text.splitlines() if line.startswith("## ")]
        missing = [s for s in REQUIRED_SECTIONS if not any(s in h for h in headings)]
        assert not missing, f"DESIGN.md 缺契约节: {missing}"

    def test_color_families_covered(self):
        text = _read()
        for family in REQUIRED_COLOR_FAMILIES:
            assert family in text, f"Colors 角色族缺失: {family}"

    def test_type_scale_categories_covered(self):
        text = _read()
        for cat in REQUIRED_TYPE_SCALE:
            assert cat in text, f"Typography 字阶类别缺失: {cat}"


class TestGuardrailEntries:
    """Do/Don't 护栏条目结构（与样式机检同源）。"""

    @staticmethod
    def _guardrail_items(text: str) -> list[str]:
        section = text.split("## Do's and Don'ts", 1)[1].split("\n## ", 1)[0]
        return [ln for ln in section.splitlines() if ln and ln[0].isdigit() and ". **DO**" in ln]

    def test_guardrail_numbering_contiguous(self):
        items = self._guardrail_items(_read())
        assert items, "护栏条目为空"
        numbers = [int(ln.split(".", 1)[0]) for ln in items]
        assert numbers == list(range(1, len(numbers) + 1)), f"护栏编号不连续: {numbers}"

    def test_guardrail_do_dont_pairs(self):
        for ln in self._guardrail_items(_read()):
            assert "**DO**" in ln and "**DON'T**" in ln, f"护栏缺 DO/DON'T 对偶: {ln[:60]}"


class TestReferencedFilesExist:
    """文内引用的实现文件真实存在（防文档腐烂）。"""

    REFS = [
        "src/static/tmpl/report_template.html",
        "src/static/web/style.css",
        "src/static/theme.js",
        "docs/managements/plan.md",
    ]

    def test_referenced_paths_exist(self):
        for rel in self.REFS:
            assert (REPO_ROOT / rel).exists(), f"DESIGN.md 引用路径不存在: {rel}"


class TestBrandMark:
    """品牌标记（猫头鹰「守望」）契约：小尺寸身份识别，不做装饰。

    契约条款见 DESIGN.md Overview「品牌标记」与护栏 11；此处把「尺寸上限」与
    「三处产物都有标记」机检化，防后续改动把它放大成 hero 或悄悄删掉。
    """

    #: 各面标记尺寸上限（px）——DESIGN.md 护栏 11
    LIMITS = {"report": 40, "web": 40, "excel": 48}

    def test_report_header_has_brand_mark(self):
        html = (REPO_ROOT / "src/static/tmpl/report_template.html").read_text(encoding="utf-8")
        assert 'class="report-brand-mark"' in html, "报告页头缺品牌标记"
        assert "width: 32px" in html, "报告标记应为 32px（≤40px 契约）"

    def test_report_favicon_is_owl_not_chart(self):
        """报告 favicon 与 Web 同源（猫头鹰头像），不得退回折线图。"""
        html = (REPO_ROOT / "src/static/tmpl/report_template.html").read_text(encoding="utf-8")
        favicon_line = next(ln for ln in html.splitlines() if 'rel="icon"' in ln)
        assert "%23e8f1fb" in favicon_line, "favicon 应为猫头鹰头像（脸盘色 #e8f1fb）"
        assert "M7 22 13.5 15.5" not in favicon_line, "favicon 不得退回折线图"

    def test_web_header_has_brand_mark(self):
        html = (REPO_ROOT / "src/static/web/index.html").read_text(encoding="utf-8")
        css = (REPO_ROOT / "src/static/web/style.css").read_text(encoding="utf-8")
        assert 'class="app-brand-mark"' in html, "工作台页头缺品牌标记"
        assert "width: 36px" in css, "工作台标记应为 36px（≤40px 契约）"

    def test_excel_summary_embeds_brand_mark(self):
        """Excel 首页标题行嵌入标记（≤48px）；资源缺失时静默跳过不阻断报告。"""
        from openpyxl import Workbook

        from src.python.report.summary import _BRAND_MARK_PNG, _BRAND_MARK_PX, _embed_brand_mark

        assert _BRAND_MARK_PNG.exists(), "品牌标记 PNG 资源缺失（应由 favicon.svg 栅格化后入库）"
        assert _BRAND_MARK_PX <= self.LIMITS["excel"], f"Excel 标记 {_BRAND_MARK_PX}px 超契约上限"
        wb = Workbook()
        ws = wb.active
        _embed_brand_mark(ws)
        assert len(ws._images) == 1, "Excel 首页应嵌入 1 个品牌标记"

    def test_brand_mark_survives_double_save(self, tmp_path):
        """报告落盘两次（最新版 + 存档版）时标记不得失效。

        回归：openpyxl 的 `Image._data()` 每次调用都 `fp.close()` 关流，第二次
        `wb.save()` 会抛 `ValueError: I/O operation on closed file`。
        """
        from openpyxl import Workbook, load_workbook

        from src.python.report.summary import _embed_brand_mark

        wb = Workbook()
        ws = wb.active
        _embed_brand_mark(ws)
        first, second = tmp_path / "latest.xlsx", tmp_path / "archive.xlsx"
        wb.save(first)
        wb.save(second)  # 不得抛异常
        assert len(load_workbook(second).active._images) == 1

    def test_brand_mark_failure_is_silent(self, monkeypatch):
        """资源缺失不得让报告生成失败（品牌标记是增强，不是依赖）——CI 无 cairo 也能出报告。"""
        from openpyxl import Workbook

        from src.python.report import summary

        monkeypatch.setattr(summary, "_BRAND_MARK_PNG", REPO_ROOT / "no-such-file.png")
        wb = Workbook()
        ws = wb.active
        summary._embed_brand_mark(ws)  # 不应抛出
        assert len(ws._images) == 0
