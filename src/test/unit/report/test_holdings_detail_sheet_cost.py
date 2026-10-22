"""持仓明细页签 — 资金加权成本与申购状态列单元测试。

从 `test_holdings_detail_sheet.py` 分出（按区块写入职责分片）：本文件只覆盖
加权成本列与申购状态列的取值口径，其余（行值映射/数字格式/着色/账户分组/
市值区块写入）留在主分片。
"""

from __future__ import annotations

import unittest

import pytest
from openpyxl import Workbook

from src.python.core.models import Holding
from src.python.report import holdings_detail_sheet as hds

pytestmark = [pytest.mark.unit, pytest.mark.unit_report, pytest.mark.usefixtures("offline_external_sources")]


class TestWeightedAvgCost(unittest.TestCase):
    """_weighted_avg_cost 单元测试：批次成本价按份额加权（含费用摊薄）。"""

    def test_returns_weighted_cost(self):
        """多档批次按份额加权 = 总成本 / 总份额。"""
        buckets = {
            "low": {"shares": 100.0, "cost": 900.0},
            "high": {"shares": 50.0, "cost": 600.0},
            "unpriced": {"shares": 0.0, "cost": 0.0},
        }
        # (900 + 600) / (100 + 50) = 1500 / 150 = 10.0
        self.assertAlmostEqual(hds._weighted_avg_cost(buckets), 10.0)

    def test_single_low_bucket(self):
        """仅低成本档时返回其平均成本价。"""
        buckets = {
            "low": {"shares": 100.0, "cost": 950.0},
            "high": {"shares": 0.0, "cost": 0.0},
            "unpriced": {"shares": 0.0, "cost": 0.0},
        }
        self.assertAlmostEqual(hds._weighted_avg_cost(buckets), 9.5)

    def test_none_for_missing_or_empty(self):
        """缺码/空桶/零份额时返回 None（渲染层写空）。"""
        self.assertIsNone(hds._weighted_avg_cost(None))
        self.assertIsNone(hds._weighted_avg_cost({}))
        self.assertIsNone(hds._weighted_avg_cost({}))
        self.assertIsNone(
            hds._weighted_avg_cost({"low": {"shares": 0.0, "cost": 0.0}, "high": {"shares": 0.0, "cost": 0.0}})
        )


# ═══════════════════════════════════════════════════════════
#  资金加权成本列渲染（成本流水子模块）
# ═══════════════════════════════════════════════════════════


class TestWriteMarketValueBlockFlow(unittest.TestCase):
    """_write_market_value_block 资金加权成本列渲染测试。

    功能开关 `cost_lots` 对应 fund_flow_data 是否传入：
    None → 保持既有 15 列；非 None → 追加「资金加权成本」列（第 16 列）。
    """

    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.detail = hds.DetailRow(
            account="证券账户",
            name="电池ETF",
            code="561910",
            price=10.0,
            nav_date="2026-06-26",
            yesterday_close=9.5,
            price_type="场内收盘价(T)",
            premium="--",
            shares=100.0,
            market_value=1000.0,
            cost=100.0,
            profit=900.0,
            profit_rate=9.0,
            today_profit=50.0,
            source="腾讯财经",
            source_api="tencent",
        )

    def _flow(self, low_shares=0.0, cost=0.0):
        return {
            "available": True,
            "cost_tiers": {
                "per_code": {
                    "561910": {
                        "low": {"shares": low_shares, "cost": cost},
                        "high": {"shares": 0.0, "cost": 0.0},
                        "unpriced": {"shares": 0.0, "cost": 0.0},
                    }
                }
            },
            "dividends": {"per_code": {}},
        }

    def test_flow_weighted_cost_header_when_enabled(self):
        """开关开启时表头第 16 列为「资金加权成本」，原 15 列保持不变。"""
        hds._write_market_value_block(self.ws, [self.detail], self._flow(), 1)
        headers = [self.ws.cell(row=2, column=c).value for c in range(1, 17)]
        self.assertEqual(len(headers), 16)
        self.assertEqual(headers[14], "取价渠道")
        self.assertEqual(headers[15], "资金加权成本")

    def test_flow_weighted_cost_value(self):
        """开关开启时数据行「资金加权成本」列 = 批次成本价按份额加权。"""
        flow = self._flow(low_shares=100.0, cost=950.0)
        hds._write_market_value_block(self.ws, [self.detail], flow, 1)
        # 数据行 row 3，列 16 = 950 / 100 = 9.5
        self.assertAlmostEqual(self.ws.cell(row=3, column=16).value, 9.5)

    def test_flow_weighted_cost_none_when_no_buckets(self):
        """开关开启但代码无批次数据时，「资金加权成本」列为空。"""
        flow = {"available": True, "cost_tiers": {"per_code": {}}, "dividends": {"per_code": {}}}
        hds._write_market_value_block(self.ws, [self.detail], flow, 1)
        self.assertIsNone(self.ws.cell(row=3, column=16).value)

    def test_no_flow_column_when_disabled(self):
        """开关关闭（fund_flow_data=None）时保持既有 15 列，无「资金加权成本」列。"""
        hds._write_market_value_block(self.ws, [self.detail], None, 1)
        headers = [self.ws.cell(row=2, column=c).value for c in range(1, 16)]
        self.assertEqual(len(headers), 15)
        self.assertEqual(headers[14], "取价渠道")
        self.assertNotIn("资金加权成本", headers)


# ═══════════════════════════════════════════════════════════
#  申购状态条件列（功能开关 fund_purchase_limit，purchase_status_data 契约）
# ═══════════════════════════════════════════════════════════


class TestPurchaseStatusColumn(unittest.TestCase):
    """purchase_status_data 契约 → 区块①「申购状态」条件列与口径脚注。

    判据单源：列可见性 = purchase_column_visible、文案 = format_purchase_status_cell、
    脚注 = purchase_status_footnote（与 HTML 端同一实现）。
    """

    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.detail = hds.DetailRow(
            account="证券账户",
            name="电池ETF",
            code="561910",
            price=10.0,
            nav_date="2026-06-26",
            yesterday_close=9.5,
            price_type="场内收盘价(T)",
            premium="--",
            shares=100.0,
            market_value=1000.0,
            cost=100.0,
            profit=900.0,
            profit_rate=9.0,
            today_profit=50.0,
            source="腾讯财经",
            source_api="tencent",
        )

    @staticmethod
    def _contract(available: bool = True) -> dict:
        return {
            "available": available,
            "reason": None if available else "全链路失败",
            "rows": {
                "561910": {
                    "purchase_status": "限大额",
                    "redemption_status": "开放赎回",
                    "next_open_date": "",
                    "daily_limit": 100.0,
                    "min_purchase": 10.0,
                }
            },
            "fetched_at": "2099-01-01T08:00:00+08:00",
            "source": "天天基金",
        }

    @staticmethod
    def _headers(ws, count: int, row: int = 2) -> list:
        return [ws.cell(row=row, column=c).value for c in range(1, count + 1)]

    def test_header_appended_when_available(self):
        """available=True → 表头第 16 列「申购状态」，原 15 列保持不变。"""
        hds._write_market_value_block(self.ws, [self.detail], None, 1, self._contract())
        headers = self._headers(self.ws, 16)
        self.assertEqual(headers[14], "取价渠道")
        self.assertEqual(headers[15], "申购状态")

    def test_cell_text_from_single_source(self):
        """数据行单元格文案与单源函数逐字一致（限大额 + 日限额）。"""
        from src.python.report.purchase_status import format_purchase_status_cell

        contract = self._contract()
        hds._write_market_value_block(self.ws, [self.detail], None, 1, contract)
        expected = format_purchase_status_cell(contract, "561910", "stale")
        self.assertEqual(self.ws.cell(row=3, column=16).value, expected)
        self.assertEqual(expected, "🟡 限大额 日限 100 元")

    def test_subtotal_and_total_cell_blank(self):
        """小计/总计行申购状态列留空（非可聚合指标）。"""
        hds._write_market_value_block(self.ws, [self.detail], None, 1, self._contract())
        for row in (4, 5):
            self.assertIn(self.ws.cell(row=row, column=16).value, (None, ""))

    def test_flow_and_purchase_coexist(self):
        """成本流水 + 申购状态同时开启 → 17 列，申购状态在末列。"""
        flow = {"available": True, "cost_tiers": {"per_code": {}}, "dividends": {"per_code": {}}}
        hds._write_market_value_block(self.ws, [self.detail], flow, 1, self._contract())
        headers = self._headers(self.ws, 17)
        self.assertEqual(len(headers), 17)
        self.assertEqual(headers[15], "资金加权成本")
        self.assertEqual(headers[16], "申购状态")
        self.assertIn("限大额", self.ws.cell(row=3, column=17).value)

    def test_hidden_when_unavailable(self):
        """available=False（全链失败）→ 静默隐列：保持既有 15 列。"""
        hds._write_market_value_block(self.ws, [self.detail], None, 1, self._contract(available=False))
        headers = self._headers(self.ws, 16)
        self.assertEqual(headers[15], None)
        self.assertNotIn("申购状态", [h for h in headers if h])

    def test_hidden_when_contract_none(self):
        """契约缺席（开关关）→ 与主干基线一致的 15 列。"""
        hds._write_market_value_block(self.ws, [self.detail], None, 1, None)
        headers = self._headers(self.ws, 15)
        self.assertEqual(len(headers), 15)
        self.assertEqual(headers[14], "取价渠道")

    def test_footnote_written_with_channel_caveat(self):
        """整页写入时区块①下方写口径脚注，文案与单源逐字一致（Excel/HTML 同源）。"""
        from src.python.report.purchase_status import PURCHASE_STATUS_FOOTNOTE, purchase_status_footnote

        contract = self._contract()
        holding = Holding(account="证券账户", name="电池ETF", code="561910", shares=100, cost_price=1.0)
        hds.write_holdings_detail_sheet(self.ws, [holding], [self.detail], None, contract)
        foot_cells = [
            c.value
            for row in self.ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and "天天基金渠道口径" in c.value
        ]
        self.assertEqual(len(foot_cells), 1, "口径脚注应恰好出现一次")
        self.assertIn(PURCHASE_STATUS_FOOTNOTE, foot_cells[0])
        self.assertEqual(foot_cells[0], purchase_status_footnote(contract))

    def test_footnote_absent_when_column_hidden(self):
        """列不渲染（契约缺席/不可用）→ 无脚注、无申购状态列（既有输出不变）。"""
        holding = Holding(account="证券账户", name="电池ETF", code="561910", shares=100, cost_price=1.0)
        hds.write_holdings_detail_sheet(self.ws, [holding], [self.detail], None, self._contract(available=False))
        values = [c.value for row in self.ws.iter_rows() for c in row if c.value is not None]
        self.assertNotIn("申购状态", values)
        self.assertFalse(any(isinstance(v, str) and "天天基金渠道口径" in v for v in values))
