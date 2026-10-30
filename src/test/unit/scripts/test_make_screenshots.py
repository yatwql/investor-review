"""`scripts/make-screenshots.py`（演示截图生成器）测试。

三块实现各测其纯逻辑面：演示工作区的持仓/配置口径与报告命令装配、TUI 帧的折行与
路径瘦身（渲染尺寸在依赖齐备时抽验）、Web 截图的浏览器解析与命令装配。
外部依赖（Pillow / CJK 字体 / Node / Chromium）缺失即跳过或走纯函数断言，不发起网络请求。
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pytest

from src.test._script_loader import load_script

pytestmark = [pytest.mark.unit, pytest.mark.unit_scripts]


@pytest.fixture(name="demo")
def fixture_demo():
    """加载演示工作区模块。"""
    return load_script("_screenshots/demo_workspace.py", module_name="_screenshots_demo_workspace_under_test")


@pytest.fixture(name="tui")
def fixture_tui():
    """加载 TUI 帧渲染模块。"""
    return load_script("_screenshots/tui_frame.py", module_name="_screenshots_tui_frame_under_test")


@pytest.fixture(name="web")
def fixture_web():
    """加载 Web 截图模块。"""
    return load_script("_screenshots/web_shot.py", module_name="_screenshots_web_shot_under_test")


@pytest.fixture(name="xlsx")
def fixture_xlsx():
    """加载 Excel 页签渲染模块。"""
    return load_script("_screenshots/xlsx_frame.py", module_name="_screenshots_xlsx_frame_under_test")


class TestDemoWorkspace:
    """演示工作区：自造持仓、绝对路径配置、报告命令。"""

    @pytest.mark.unit_scripts
    def test_builds_one_sheet_per_account_with_declared_rows(self, demo, tmp_path):
        workspace = demo.build(tmp_path / "shots")
        assert (workspace / "data" / "holdings" / demo.HOLDINGS_FILENAME).is_file()

        from openpyxl import load_workbook

        workbook = load_workbook(workspace / "data" / "holdings" / demo.HOLDINGS_FILENAME)
        accounts = list(dict.fromkeys(row[0] for row in demo.HOLDINGS))
        assert workbook.sheetnames == accounts
        for account in accounts:
            rows = list(workbook[account].iter_rows(values_only=True))
            assert rows[0] == ("名称", "代码", "持仓份额", "每份成本")
            expected = [row[1:] for row in demo.HOLDINGS if row[0] == account]
            assert [tuple(row) for row in rows[1:]] == expected

    @pytest.mark.unit_scripts
    def test_config_uses_absolute_paths(self, demo, tmp_path):
        workspace = demo.build(tmp_path / "shots")
        config = json.loads((workspace / "data" / "config" / "config.json").read_text(encoding="utf-8"))

        assert Path(config["holdings_dir"]).is_absolute()
        assert Path(config["output_dir"]).is_absolute()
        assert not (workspace / "display-config.json").exists()

    @pytest.mark.unit_scripts
    def test_placeholder_key_is_obviously_fake(self, demo, tmp_path):
        workspace = demo.build(tmp_path / "shots")
        key = json.loads((workspace / "data" / "config" / "llm_key.json").read_text(encoding="utf-8"))
        block = key[demo.DEMO_CREDENTIALS_REF]
        assert block["api_key"] == demo.DEMO_API_KEY
        assert "demo" in block["api_key"]

    @pytest.mark.unit_scripts
    def test_demo_credentials_match_provider_refs(self, demo, tmp_path):
        """凭据块名与 providers 的 credentials_ref 必须同名。

        回归场景：demo 的 llm_key.json 只有一个旧格式的 ``default`` 块，而 providers 引用
        ``deepseek-main`` / ``opencode-go``——每次演示 run 刷两行「引用凭据不存在」告警，
        污染真实日志的告警统计。
        """
        workspace = demo.build(tmp_path / "shots")
        key = json.loads((workspace / "data" / "config" / "llm_key.json").read_text(encoding="utf-8"))
        providers = json.loads((workspace / "data" / "config" / "llm_providers.json").read_text(encoding="utf-8"))
        refs = {p["credentials_ref"] for p in providers["providers"]}
        assert refs == set(key), "演示工作区的凭据块与 provider 引用必须一一对应"

    @pytest.mark.unit_scripts
    def test_isolation_env_redirects_logs_and_cache(self, demo, tmp_path):
        """演示 run 的日志/缓存必须落演示目录，真实 logs//data/cache 零接触。"""
        workspace = demo.build(tmp_path / "shots")
        env = demo.isolation_env(workspace)
        assert env["INVEST_RUN_CONTEXT"] == demo.DEMO_RUN_CONTEXT
        assert env["INVEST_LOG_PATH"].startswith(str(workspace))
        assert env["INVEST_CACHE_DIR"].startswith(str(workspace))
        assert str(demo.REPO_ROOT / "logs" / "app.log") != env["INVEST_LOG_PATH"]

    @pytest.mark.unit_scripts
    def test_report_command_runs_report_subcommand_on_demo_config(self, demo, tmp_path):
        config = demo.config_path(tmp_path / "shots")
        command = demo.report_command(config)
        assert command[0] == sys.executable
        assert command[1:3] == ["-m", "src.python.cli"]
        assert command[command.index("--config") + 1] == str(config)
        assert "report" in command
        assert command[-1] == "--yes"


class TestCliText:
    """CLI 输出清洗：剥时间戳、进度回车只留末段、抹绝对路径、截尾。"""

    @pytest.mark.unit_scripts
    def test_strips_log_stamp_and_keeps_info_prefix(self, demo):
        text = "2026-10-27 07:52:43,331 [INFO] 总市值: 1.00元\n"
        assert demo.clean_cli_text(text) == ["[INFO] 总市值: 1.00元"]

    @pytest.mark.unit_scripts
    def test_progress_carriage_returns_keep_last_state(self, demo):
        text = "进度 10%\r进度 60%\r进度 100%\n"
        assert demo.clean_cli_text(text) == ["进度 100%"]

    @pytest.mark.unit_scripts
    def test_scrubs_repo_root_to_relative(self, demo):

        line = f"[INFO] 耗时已记录到 {demo.REPO_ROOT}/data/state/perf_history.jsonl"
        assert demo.clean_cli_text(line) == ["[INFO] 耗时已记录到 data/state/perf_history.jsonl"]

    @pytest.mark.unit_scripts
    def test_crlf_lines_are_kept(self, demo):
        """回归：终端输出是 CRLF，先 split("\\r") 会把整行当空段丢掉（实测 23KB 只剩 1 行）。"""
        text = "2026-10-27 07:52:43,331 [INFO] 第一行\r\n2026-10-27 07:52:43,332 [INFO] 第二行\r\n"
        assert demo.clean_cli_text(text) == ["[INFO] 第一行", "[INFO] 第二行"]

    @pytest.mark.unit_scripts
    def test_tail_window_keeps_last_lines(self, demo):
        text = "\n".join(f"[INFO] 行{i}" for i in range(30))
        lines = demo.clean_cli_text(text, tail_lines=5)
        assert lines == [f"[INFO] 行{i}" for i in range(25, 30)]


class TestStateGuard:
    """演示报告不得改动真实运行状态：跑前快照、跑后还原（含删除运行期新增文件）。"""

    @pytest.mark.unit_scripts
    def test_snapshot_and_restore_roundtrip(self, tmp_path, demo):
        state = tmp_path / "state"
        (state / "nested").mkdir(parents=True)
        (state / "perf.jsonl").write_text("old\n", encoding="utf-8")
        (state / "nested/k.json").write_text("{}", encoding="utf-8")

        snap = demo.snapshot_paths((state,))
        assert set(snap) == {state / "perf.jsonl", state / "nested/k.json"}

        (state / "perf.jsonl").write_text("new\n", encoding="utf-8")
        (state / "added.json").write_text("x", encoding="utf-8")
        demo.restore_paths((state,), snap)

        assert (state / "perf.jsonl").read_text(encoding="utf-8") == "old\n"
        assert not (state / "added.json").exists(), "运行期新增文件应被删除"

    @pytest.mark.unit_scripts
    def test_snapshot_missing_dir_is_noop(self, tmp_path, demo):
        assert demo.snapshot_paths((tmp_path / "absent",)) == {}
        demo.restore_paths((tmp_path / "absent",), {})


class TestTuiFrame:
    """TUI 帧：折行保色、路径瘦身、渲染尺寸。"""

    @pytest.mark.unit_scripts
    def test_short_line_is_returned_untouched(self, tui):
        line = "\x1b[1m\x1b[32m首页\x1b[0m"
        assert tui.wrap_line(line, 40) == [line]

    @pytest.mark.unit_scripts
    def test_wrapped_rows_respect_width_and_reopen_color(self, tui):
        plain = "中" * 60
        rows = tui.wrap_line(f"\x1b[32m{plain}\x1b[0m", 20)
        assert len(rows) > 1
        assert [tui.strip_ansi(row) for row in rows] == [plain[index : index + 10] for index in range(0, 60, 10)]
        for row in rows:
            assert tui.display_width(tui.strip_ansi(row)) <= 20
        # 续行重新打开配色，首行带原始起始序列
        assert rows[0].startswith("\x1b[32m")
        assert all(row.startswith("\x1b[32m") for row in rows[1:])

    @pytest.mark.unit_scripts
    def test_scrub_paths_removes_longest_prefix_first(self, tui, tmp_path):
        demo_dir = tmp_path / "shots" / "demo"
        text = f"{demo_dir}/reports/报告.html 与 {tmp_path}/README.md"
        scrubbed = tui.scrub_paths(text, (demo_dir, tmp_path))
        assert str(demo_dir) not in scrubbed
        assert str(tmp_path) not in scrubbed
        assert "reports/报告.html" in scrubbed
        assert "README.md" in scrubbed

    @pytest.mark.unit_scripts
    def test_render_sizes_canvas_from_line_count(self, tui, tmp_path):
        pytest.importorskip("PIL.Image")
        try:
            tui.resolve_font()
        except tui.RenderDependencyError as exc:
            pytest.skip(str(exc))
        output = tui.render("\x1b[32m第一行\x1b[0m\n第二行\n", tmp_path / "frame.png")
        from PIL import Image

        with Image.open(output) as image:
            assert image.size[1] == tui.LINE_HEIGHT * 2 + tui.PADDING * 2
            assert image.size[0] >= 320

    @pytest.mark.unit_scripts
    def test_colored_tail_starts_at_measured_offset(self, tui, tmp_path):
        """回归：配色段起点按字体实测推进。

        旧实现按「字号 × 固定系数」估算每段步进，等宽 CJK 字体里拉丁/全角的实际
        步进与估算不一致，偏差随段落长度累积——页头状态行这类长行上表现为彩色段
        前的一大段空白（README 配图曾如此）。
        """
        pytest.importorskip("PIL.Image")
        from PIL import Image, ImageDraw, ImageFont

        try:
            font_path = tui.resolve_font()
        except tui.RenderDependencyError as exc:
            pytest.skip(str(exc))
        prefix = "中文字符" * 3
        output = tui.render(f"{prefix}\x1b[32m末段\x1b[0m\n", tmp_path / "drift.png")
        font = ImageFont.truetype(font_path, tui.FONT_SIZE)
        probe = ImageDraw.Draw(Image.new("RGB", (1, 1)))
        expected_x = tui.PADDING + probe.textlength(prefix, font=font)
        green = tui.PALETTE[32]
        hits: list[int] = []
        with Image.open(output) as image:
            pixels = image.convert("RGB").load()
            assert pixels is not None
            for y in range(tui.PADDING, tui.PADDING + tui.LINE_HEIGHT):
                for x in range(image.size[0]):
                    if pixels[x, y] == green:
                        hits.append(x)
        assert hits, "未找到绿色末段的像素"
        assert abs(min(hits) - expected_x) <= 4, f"配色段偏移 {min(hits) - expected_x:.1f}px（估算步进残留）"


class TestWebShot:
    """Web 截图：浏览器解析次序与命令装配。"""

    @pytest.mark.unit_scripts
    def test_explicit_chrome_wins_over_env(self, web, tmp_path, monkeypatch):
        explicit = tmp_path / "chrome"
        fallback = tmp_path / "fallback"
        for path in (explicit, fallback):
            path.write_text("", encoding="utf-8")
        monkeypatch.setenv("CHROME_BIN", str(fallback))
        assert web.resolve_chrome(str(explicit)) == explicit

    @pytest.mark.unit_scripts
    def test_missing_explicit_falls_back_to_env(self, web, tmp_path, monkeypatch):
        fallback = tmp_path / "fallback"
        fallback.write_text("", encoding="utf-8")
        monkeypatch.setenv("CHROME_BIN", str(fallback))
        assert web.resolve_chrome(str(tmp_path / "absent")) == fallback

    @pytest.mark.unit_scripts
    def test_build_command_maps_positional_arguments(self, web, tmp_path):
        command = web.build_command(
            "http://127.0.0.1:8000/workbench",
            tmp_path / "shot.png",
            scroll_y=120,
            width=1280,
            height=800,
            replace="临时=>演示",
        )
        assert Path(command[1]).name == "cdp_shot.mjs"
        assert command[2:] == [
            "http://127.0.0.1:8000/workbench",
            str(tmp_path / "shot.png"),
            "120",
            "1280",
            "800",
            "临时=>演示",
        ]

    @pytest.mark.unit_scripts
    def test_cdp_script_reads_chrome_from_environment(self, web):
        source = web.CDP_SCRIPT.read_text(encoding="utf-8")
        assert "CHROME_BIN" in source
        assert "remote-debugging-port" in source


class TestXlsxFrame:
    """Excel 页签 → 表格图：显示值换算、超链接取文本、合并区与列宽收敛。"""

    @pytest.mark.unit_scripts
    def test_number_format_follows_cell_display_rules(self, xlsx):
        """百分比/千分位/通用格式各自按 Excel 显示口径收敛。"""
        assert xlsx.format_value(-0.0286, "0.00%") == "-2.86%"
        assert xlsx.format_value(368152.5, "#,##0.00") == "368,152.50"
        assert xlsx.format_value(0.0, "0.00%") == "0.00%"
        assert xlsx.format_value(6, "General") == "6"
        assert xlsx.format_value(4.065, "General") == "4.065"
        assert xlsx.format_value(None, "General") == ""
        assert xlsx.format_value("--", "General") == "--"

    @pytest.mark.unit_scripts
    def test_hyperlink_formula_renders_its_display_text(self, xlsx):
        """页签导航整行是 HYPERLINK 公式：上屏取第二个参数，不是公式原文。"""
        value = '=HYPERLINK("#\'1.投资分析汇总\'!A1","1.投资分析汇总")'
        assert xlsx.hyperlink_text(value) == "1.投资分析汇总"
        assert xlsx.format_value(value, "General") == "1.投资分析汇总"
        assert xlsx.hyperlink_text("=SUM(A1:A2)") is None
        assert xlsx.hyperlink_text(3) is None

    @pytest.mark.unit_scripts
    def test_column_letter_extends_past_single_letter(self, xlsx):
        """列标 A→Z→AA 连续，超过 26 列不失序。"""
        assert xlsx.column_letter(1) == "A"
        assert xlsx.column_letter(26) == "Z"
        assert xlsx.column_letter(27) == "AA"

    @pytest.mark.unit_scripts
    def test_html_keeps_merges_tabs_and_escapes_content(self, xlsx, tmp_path):
        """合并区出 colspan、页签条列出全部页签且标出当前页、单元格文本被转义。"""
        from openpyxl import Workbook

        workbook = Workbook()
        first = workbook.worksheets[0]
        first.title = "汇总"
        first["A1"] = "<b>标题</b>"
        first.merge_cells("A1:B1")
        first["A2"] = 0.1234
        first["A2"].number_format = "0.00%"
        second = workbook.create_sheet("明细")
        second["A1"] = 1
        path = tmp_path / "报告.xlsx"
        workbook.save(path)

        markup = xlsx.build_html(path, viewport_width=1200)

        assert 'colspan="2"' in markup
        assert "&lt;b&gt;标题&lt;/b&gt;" in markup
        assert "12.34%" in markup
        assert 'class="tab active">汇总<' in markup
        assert '<span class="tab">明细</span>' in markup

    @pytest.mark.unit_scripts
    def test_wide_sheet_shrinks_to_viewport(self, xlsx, tmp_path):
        """列总宽超视口时等比收缩，保证整行不被截断（列标数仍等于请求列数）。"""
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "宽表"
        for col in range(1, 21):
            sheet.cell(row=1, column=col, value="很长的中文说明文本" * 2)
        path = tmp_path / "宽表.xlsx"
        workbook.save(path)

        markup = xlsx.build_html(path, max_cols=20, viewport_width=1200)
        widths = [int(value) for value in re.findall(r'col style="width:(\d+)px"', markup)]

        assert len(widths) == 20
        assert sum(widths) <= 1200 - 2 * 24 - xlsx._ROWNO_PX
        assert sum(1 for _ in markup.split('<th class="colno">')) - 1 == 20


class TestEntry:
    """入口：子命令分发与参数错误口径。"""

    @pytest.mark.unit_scripts
    def test_demo_subcommand_writes_workspace(self, tmp_path, monkeypatch):
        entry = load_script("make-screenshots.py", module_name="_make_screenshots_under_test")
        monkeypatch.setattr(entry.demo_workspace, "DEFAULT_ROOT", tmp_path / "shots")
        assert entry.main(["demo"]) == 0
        holdings = tmp_path / "shots" / "demo" / "data" / "holdings"
        assert (holdings / entry.demo_workspace.HOLDINGS_FILENAME).is_file()

    @pytest.mark.unit_scripts
    def test_report_subcommand_returns_underlying_exit_code(self, tmp_path, monkeypatch):
        entry = load_script("make-screenshots.py", module_name="_make_screenshots_under_test")
        # 入口的演示持仓护栏读 DEFAULT_ROOT：必须指到 tmp 并真建工作区，
        # 否则用例会依赖本地 git 忽略的 docs/tmp/screenshots 产物（CI 上必红）
        monkeypatch.setattr(entry.demo_workspace, "DEFAULT_ROOT", tmp_path / "shots")
        assert entry.main(["demo"]) == 0
        monkeypatch.setattr(entry.demo_workspace, "run_report", lambda root=None: 3)
        assert entry.main(["report"]) == 3

    @pytest.mark.unit_scripts
    def test_report_subcommand_refuses_without_demo_holdings(self, tmp_path, monkeypatch):
        """演示工作区缺失时入口护栏直接拒绝，不落到报告生成（零真实数据接触）。"""
        entry = load_script("make-screenshots.py", module_name="_make_screenshots_under_test")
        monkeypatch.setattr(entry.demo_workspace, "DEFAULT_ROOT", tmp_path / "shots")
        called: list[int] = []
        monkeypatch.setattr(entry.demo_workspace, "run_report", lambda root=None: called.append(1) or 0)
        with pytest.raises(RuntimeError):
            entry.main(["report"])
        assert called == []

    @pytest.mark.unit_scripts
    def test_unknown_subcommand_exits_with_usage_error(self):
        entry = load_script("make-screenshots.py", module_name="_make_screenshots_under_test")
        with pytest.raises(SystemExit) as caught:
            entry.main(["bogus"])
        assert caught.value.code == 2

    @pytest.mark.unit_scripts
    def test_help_lists_every_subcommand(self, capsys):
        entry = load_script("make-screenshots.py", module_name="_make_screenshots_under_test")
        with pytest.raises(SystemExit) as caught:
            entry.main(["--help"])
        assert caught.value.code == 0
        printed = capsys.readouterr().out
        assert {"demo", "report", "tui", "web", "xlsx"} <= set(printed.split())

    @pytest.mark.unit_scripts
    def test_xlsx_rows_default_comes_from_renderer_constant(self, xlsx, tmp_path):
        """`--rows` 缺省值取自渲染模块常量（行数口径只有一个来源）。"""
        entry = load_script("make-screenshots.py", module_name="_make_screenshots_under_test")
        parser = entry.build_parser()
        parsed = parser.parse_args(["xlsx", str(tmp_path / "a.xlsx"), str(tmp_path / "a.png")])
        assert parsed.rows == xlsx.DEFAULT_MAX_ROWS

    @pytest.mark.unit_scripts
    def test_xlsx_subcommand_forwards_sheet_and_size(self, tmp_path, monkeypatch):
        """`xlsx` 子命令把页签名与目标尺寸原样传下去（护栏不过则不出图）。"""
        entry = load_script("make-screenshots.py", module_name="_make_screenshots_under_test")
        monkeypatch.setattr(entry.demo_workspace, "DEFAULT_ROOT", tmp_path / "shots")
        assert entry.main(["demo"]) == 0
        captured: dict[str, object] = {}

        def fake_render(workbook, output, **kwargs):
            captured.update({"workbook": workbook, "output": output, **kwargs})
            return 0

        monkeypatch.setattr(entry.xlsx_frame, "render", fake_render)
        workbook = tmp_path / "报告.xlsx"
        workbook.write_bytes(b"")
        assert entry.main(["xlsx", str(workbook), str(tmp_path / "out.png"), "--sheet", "明细", "--rows", "9"]) == 0
        assert captured["sheet_name"] == "明细"
        assert captured["max_rows"] == 9
        assert captured["output"] == tmp_path / "out.png"
